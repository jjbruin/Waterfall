"""The rent roll exhibit for the investment committee memo -- step 5 of the plan.

New business's template: `Claude - Market at Poplar Rent Roll.xlsx`, range B2:H111
(Sep 29 2026). The presentation is part of the requirement ("pay attention to the
formatting, shading, and lines"), so every property below was read off that file,
not approximated: see `.claude/memory/rent_roll_exhibit.md`, "The exhibit, exactly".

NOTHING IS COMPUTED HERE. Every figure comes from `get_rent_roll_timeline` -- the
one timeline engine, with the analyst's settlements applied -- so the exhibit
cannot differ from the Rent Roll screen. This module only lays it out.

Paging: the header row repeats on every printed page (Excel's print titles) and a
page never breaks inside a tenant's group. Their template achieved the same by
pasting a second header row at row 56; print titles give the printed result
without a header row sitting in the middle of the data.
"""
from __future__ import annotations

import io
from datetime import date
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color
from openpyxl.worksheet.pagebreak import Break

FONT = 'Times New Roman'
SIZE = 12
#: Theme colour 0 (white) darkened 5% -- the template's header and total fill.
GREY = PatternFill(fill_type='solid', fgColor=Color(theme=0, tint=-0.0499893185216834))
MEDIUM = Side(style='medium')
THIN = Side(style='thin')
WIDTHS = {'A': 2.0, 'B': 40.0, 'C': 10.0, 'D': 13.54, 'F': 14.0, 'G': 12.0, 'H': 23.45}
HEADERS = ['Tenant', 'SF', 'Start', 'Expiration', 'Annual Rent', 'Annual PSF', 'Option(s)']
FIRST_COL, LAST_COL = 2, 8          # B .. H
#: Rows per printed page before a group must move to the next page -- the
#: template's first page runs from row 3 to row 55.
ROWS_PER_PAGE = 53

BASIS_TEXT = {'fmv': 'Fair market', 'percent_of_fmv': '% of fair market',
              'greater_of': 'Greater of', 'cpi': 'CPI', 'not_stated': 'Not stated',
              'stated_increase': 'Stated increase'}


def _d(v) -> Optional[date]:
    if not v:
        return None
    try:
        return date.fromisoformat(str(v)[:10])
    except ValueError:
        return None


def _font(bold=False, italic=False) -> Font:
    return Font(name=FONT, size=SIZE, bold=bold, italic=italic)


def _frame(ws, r: int, bottom: Optional[Side] = None, top: Optional[Side] = None) -> None:
    """Medium frame left of B and right of H on every row; an optional rule."""
    for c in range(FIRST_COL, LAST_COL + 1):
        cell = ws.cell(r, c)
        if cell.value is None:
            cell.font = _font()      # Times New Roman throughout, blank cells too
        cell.border = Border(left=MEDIUM if c == FIRST_COL else None,
                             right=MEDIUM if c == LAST_COL else None,
                             top=top, bottom=bottom)


def _header(ws, r: int) -> None:
    for i, h in enumerate(HEADERS):
        cell = ws.cell(r, FIRST_COL + i, h)
        cell.font = _font(bold=True)
        cell.fill = GREY
        cell.alignment = Alignment(horizontal='center')
    _frame(ws, r, bottom=THIN, top=MEDIUM)


def _row(ws, r: int, row: Dict[str, Any], tenant_name: str, basis: Optional[str]) -> None:
    kind = row['kind']
    right = Alignment(horizontal='right')
    if kind == 'tenant':
        ws.cell(r, 2, tenant_name).font = _font()
        ws.cell(r, 2).alignment = Alignment(horizontal='left')
        c = ws.cell(r, 3, row.get('sf'))
        c.number_format, c.font, c.alignment = '#,##0', _font(), right
        o = ws.cell(r, 8, row.get('options') or 'None')
        o.font, o.alignment = _font(), right
    else:
        lab = ws.cell(r, 2, row.get('label'))
        lab.font, lab.alignment = _font(italic=True), right
    for col, key in ((4, 'start'), (5, 'end')):
        c = ws.cell(r, col, _d(row.get(key)))
        c.number_format, c.font, c.alignment = 'mm-dd-yy', _font(), right
    rent = row.get('annual_rent')
    if rent is None and kind != 'tenant':
        c = ws.cell(r, 6, BASIS_TEXT.get((basis or 'not_stated').lower(), 'Not stated'))
        c.font, c.alignment = _font(italic=True), right
    else:
        c = ws.cell(r, 6, rent)
        c.number_format, c.font, c.alignment = '#,##0', _font(), right
    c = ws.cell(r, 7, row.get('psf'))
    c.number_format, c.font, c.alignment = '#,##0.00', _font(), right


def build_exhibit(rent_roll: Dict[str, Any], title: Optional[str] = None) -> bytes:
    """The workbook: sheet 1 is the exhibit (B2:H{n}); sheet 2 lists every flag."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Rent Roll'
    ws.sheet_view.showGridLines = False
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    _header(ws, 2)
    r = 3
    page_start = 3
    breaks: List[int] = []
    for t in rent_roll['tenants']:
        rows = t['rows']
        if r - page_start + len(rows) > ROWS_PER_PAGE and r > page_start:
            breaks.append(r - 1)
            page_start = r
        opt_basis = {o['label']: o.get('rent_basis') for o in t.get('options') or []}
        current_opt = None
        for i, row in enumerate(rows):
            if row['kind'] == 'option':
                current_opt = row.get('label')
            _row(ws, r, row, t.get('tenant_name') or '', opt_basis.get(current_opt))
            _frame(ws, r, bottom=THIN if i == len(rows) - 1 else None)
            r += 1
    totals = rent_roll['totals']
    lab = ws.cell(r, 2, 'Total / Wtd. Avg.')
    for c, val, fmt in ((3, totals.get('square_feet'), '#,##0'),
                        (6, totals.get('annual_rent'), '#,##0'),
                        (7, totals.get('psf'), '#,##0.00')):
        cell = ws.cell(r, c, val)
        cell.number_format = fmt
    for c in range(FIRST_COL, LAST_COL + 1):
        cell = ws.cell(r, c)
        cell.font = _font(bold=True)
        cell.fill = GREY
        if c != FIRST_COL:
            cell.alignment = Alignment(horizontal='right')
    lab.alignment = Alignment(horizontal='left')
    _frame(ws, r, bottom=MEDIUM)
    ws.print_title_rows = '2:2'
    ws.print_area = 'B2:H%d' % r
    for b in breaks:
        ws.row_breaks.append(Break(id=b))
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if title:
        ws.oddHeader.center.text = title

    fl = wb.create_sheet('Flags')
    fl.append(['Tenant', 'Flag', 'Detail'])
    for c in fl[1]:
        c.font = Font(bold=True)
    for t in rent_roll['tenants']:
        for f in t.get('flags') or []:
            fl.append([t.get('tenant_name'), f.get('code', '').replace('_', ' '),
                       f.get('message')])
    fl.column_dimensions['A'].width = 34
    fl.column_dimensions['B'].width = 26
    fl.column_dimensions['C'].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
