"""Guardrail: accounting's allocation overrides in the ownership walk (Oct 5 2026).

Accounting's BRNERD / JBFAIR / NOTTNV exceptions: MRI's commitment ratios are
ENTITY-level, so PSC3's ratios apply to everything PSC3 holds. An override says
"for THIS investment, from THIS date, PSC3's investors funded THESE amounts".

Built on pe_exposure_check's fixture, plus a second deal through PSC3. Pinned:
  1. The override replaces the entity's split for its investment, entered as
     AMOUNTS (share = amount / total), and the walk returns to commitments above.
  2. ONLY that investment: another deal through the same PSC3 is untouched.
  3. ONLY from the effective date: an earlier quarter is untouched; the latest
     set effective on or before the date wins; removing a set restores the one
     before, and the record is kept.
  4. It can sit at the holding entity itself (Brainerd's PPIBPA).
  5. The report says so -- the route names it, the row lists it, a note
     explains it, the workbook's routes sheet carries it -- and the report cache
     does not serve the split from before the change.
  6. What cannot be true is refused with the reason; an investor MRI does not
     carry for that entity is saved with a warning.
  7. Only accounting writes -- the app's one rule, has_accounting_authority: an
     accounting-capable role AND the Accounting section (Jim, Oct 2 2026). An
     analyst is refused; the admin ROLE without Accounting is refused; everyone
     can read.

Usage: python scripts/allocation_override_check.py [--inject=ignore|nodate|global]
"""
import io
import os
import sys
import tempfile
from datetime import date, datetime, timedelta, timezone

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pandas as pd  # noqa: E402

import pe_exposure_check as base  # noqa: E402

INJECT = next((a.split("=", 1)[1] for a in sys.argv[1:] if a.startswith("--inject=")), "")
_passed, _failed = [], []


def chk(label, cond, detail=None):
    (_passed if cond else _failed).append(label)
    print(("   ok   " if cond else "   FAIL ") + label + ("" if cond or detail is None else "   [%s]" % (detail,)))


def near(a, b, tol=1e-9):
    return a is not None and b is not None and abs(float(a) - float(b)) <= tol


Q1, Q2, MAY = date(2026, 3, 31), date(2026, 6, 30), date(2026, 5, 15)

# A second investment through the same PSC3, to prove the override is per investment.
base.DEALS.append(("P8", "DEALE", "Echo Commons", "USD"))
base.COMMITMENTS += [("DEALE", "PPIE", 200, "2022-01-01", None),
                     ("PPIE", "PSC3", 100, "2022-01-01", None)]
base.ACCT.append(("DEALE", "PPIE", "2022-06-01", "Contribution", "Contribution: Investments", -200, "Y"))


def main() -> int:
    tmp = tempfile.mkdtemp(prefix="alloc_override_check_")
    db = os.path.join(tmp, "check.db")
    os.environ.pop("DATABASE_URL", None)
    os.environ["DB_PATH"] = db
    eng = base.build_db(db)
    import jwt
    from flask_app import create_app
    from flask_app.auth.models import create_user, list_users
    app = create_app()
    app.config["DATABASE_URL"] = None
    from flask_app.services import pe_exposure_service as pe
    from flask_app.services import allocation_override_service as aos
    from flask_app.services import ownership_chain_service as oc

    if INJECT == "ignore":
        aos.in_effect = lambda as_of, engine=None: {}
    if INJECT == "nodate":
        _orig = aos.list_overrides
        def _all_dated_early(engine=None, include_deleted=True):
            return [{**o, "effective_date": "1900-01-01"} for o in _orig(engine, include_deleted)]
        aos.list_overrides = _all_dated_early
    if INJECT == "global":
        _orig_ie = aos.in_effect
        def _any_investment(as_of, engine=None):
            got = _orig_ie(as_of, engine)
            return {(e, inv2): o for (e, _), o in got.items() for inv2 in ("DEALC", "DEALE", "DEALA")}
        aos.in_effect = _any_investment

    data = {"acct": pd.read_sql("SELECT * FROM accounting", eng), "inv": pd.read_sql("SELECT * FROM deals", eng),
            "wf": pd.read_sql("SELECT * FROM waterfalls", eng),
            "commitments_raw": pd.read_sql("SELECT * FROM commitments", eng)}

    def row(rep, inv, holder):
        return next((r for r in rep["rows"] if r["investment_id"] == inv and r["holder"] == holder), None)

    with app.app_context():
        before = pe.get_report(Q2, data=data, engine=eng)
        c0 = row(before, "DEALC", "PPIC")
        chk("before: DEALC through PSC3 is PSC 60 / Declaration 30 / F&F 10",
            c0 and near(c0["shares"]["PSC"], .6) and near(c0["shares"]["Declaration"], .3), c0 and c0["shares"])

        print("\n6. What cannot be true is refused")
        bad = [
            (("PSC3", "DEALA", "2026-04-01", [{"investor_id": "OWPSC", "amount": 1}], "x"), "not in DEALA's ownership chain"),
            (("OWPSC", "DEALC", "2026-04-01", [{"investor_id": "X", "amount": 1}], "x"), "named holder"),
            (("PSC3", "DEALC", "2026-04-01", [{"investor_id": "OWPSC", "amount": -5}], "x"), "negative"),
            (("PSC3", "DEALC", "2026-04-01", [{"investor_id": "OWPSC", "amount": 0}], "x"), "total nothing"),
            (("PSC3", "DEALC", "2026-04-01", [{"investor_id": "OWPSC", "amount": 1},
                                              {"investor_id": "owpsc", "amount": 2}], "x"), "named twice"),
            (("PSC3", "DEALC", "2026-04-01", [{"investor_id": "OWPSC", "amount": 1}], "  "), "Say why"),
            (("PSC3", "DEALC", "4/1/2026", [{"investor_id": "OWPSC", "amount": 1}], "x"), "must be a date"),
            (("PSC3", "DEALC", "2026-04-01", [{"investor_id": "OWPSC", "amount": "abc"}], "x"), "must be a number"),
        ]
        for args, why in bad:
            try:
                aos.create(*args, actor="t", engine=eng)
                chk("refused: %s" % why, False, "accepted")
            except ValueError as e:
                chk("refused: %s" % why, why in str(e), str(e))

        print("\n1-2. Replaces PSC3's split for DEALC only, as amounts")
        o1 = aos.create("psc3", "dealc", "2026-04-01",
                        [{"investor_id": "OWPSC", "amount": "900,000"}, {"investor_id": "DCXVIA", "amount": 0},
                         {"investor_id": "BRECO", "amount": 100000}],
                        "DCXVIA opted out of DEALC; its commitment was returned", "acct", engine=eng)
        chk("amounts become shares: 900,000 / 1,000,000 = 90%",
            near(next(l for l in o1["lines"] if l["investor_id"] == "OWPSC")["pct"], 90.0), o1["lines"])
        chk("an investor MRI carries for PSC3 raises no warning", not o1["warnings"], o1["warnings"])
        r2 = pe.get_report(Q2, data=data, engine=eng)
        c = row(r2, "DEALC", "PPIC")
        chk("the cache does not serve the split from before the override",
            c and not near(c["shares"]["PSC"], .6), c and c["shares"])
        chk("DEALC at Q2: PSC 90 / Declaration 0 / F&F 10",
            c and near(c["shares"]["PSC"], .9) and near(c["shares"]["Declaration"], 0)
            and near(c["shares"]["F&F"], .1), c and c["shares"])
        chk("...shares still total 100%", c and near(c["share_total"], 1.0))
        e = row(r2, "DEALE", "PPIE")
        chk("DEALE, through the SAME PSC3, is untouched (60/30/10)",
            e and near(e["shares"]["PSC"], .6) and near(e["shares"]["Declaration"], .3), e and e["shares"])
        a = row(r2, "DEALA", "PPIA")
        chk("DEALA, which never reaches PSC3, is untouched",
            a and not a["overrides"] and near(a["shares"]["PSC"], 500 / 1200 + 700 / 1200 * .15), a and a["shares"])

        print("\n5. The report says so")
        chk("every route through PSC3 names the override",
            c and all(rt["overrides"] and rt["overrides"][0]["entity"] == "PSC3" for rt in c["routes"]),
            c and c["routes"])
        chk("the row lists it", c and len(c["overrides"]) == 1 and c["overrides"][0]["effective_date"] == "2026-04-01")
        chk("a note explains it", any("PSC3" in n and "allocation override" in n and "opted out" in n
                                      for n in r2["notes"]), r2["notes"])
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(pe.to_excel(r2)))
        ws = wb["Ownership Routes"]
        txt = [str(cell.value or "") for row_ in ws.iter_rows() for cell in row_]
        chk("the workbook's routes sheet carries it", any("PSC3 from 2026-04-01" in t for t in txt))

        print("\n3. Only from the effective date; the latest set in force wins")
        r1 = pe.get_report(Q1, data=data, engine=eng)
        c1 = row(r1, "DEALC", "PPIC")
        chk("Q1 (before 4/1) is untouched: 60/30/10",
            c1 and near(c1["shares"]["PSC"], .6) and not c1["overrides"], c1 and c1["shares"])
        o2 = aos.create("PSC3", "DEALC", "2026-06-01",
                        [{"investor_id": "OWPSC", "amount": 500}, {"investor_id": "DCXVIA", "amount": 500}],
                        "a later revision", "acct", engine=eng)
        chk("at Q2 the 6/1 set wins: PSC 50 / Declaration 50",
            near(row(pe.get_report(Q2, data=data, engine=eng), "DEALC", "PPIC")["shares"]["PSC"], .5))
        chk("on 5/15 the 4/1 set still applies: PSC 90",
            near(row(pe.build(MAY, data=data, engine=eng), "DEALC", "PPIC")["shares"]["PSC"], .9))
        try:
            aos.create("PSC3", "DEALC", "2026-06-01", [{"investor_id": "OWPSC", "amount": 1}], "dup", "acct", engine=eng)
            chk("a second set on the same date is refused", False)
        except ValueError as ex:
            chk("a second set on the same date is refused", "remove it first" in str(ex), str(ex))
        aos.remove(o2["id"], "acct2", engine=eng)
        chk("removing the 6/1 set restores the 4/1 one at Q2",
            near(row(pe.get_report(Q2, data=data, engine=eng), "DEALC", "PPIC")["shares"]["PSC"], .9))
        kept = next(o for o in aos.list_overrides(eng) if o["id"] == o2["id"])
        chk("...and the removed set is kept, saying who removed it", kept["deleted_by"] == "acct2")

        print("\n4. At the holding entity itself (Brainerd's PPIBPA)")
        o3 = aos.create("PPIC", "DEALC", "2026-06-15",
                        [{"investor_id": "PSC3", "amount": 50}, {"investor_id": "NEWCO", "amount": 50}],
                        "restructure", "acct", engine=eng)
        chk("an investor MRI does not carry for the entity is saved WITH a warning",
            o3["warnings"] and "NEWCO" in o3["warnings"][0], o3["warnings"])
        c3 = row(pe.get_report(Q2, data=data, engine=eng), "DEALC", "PPIC")
        chk("the root's split is replaced and PSC3's override still applies below it: PSC 45, F&F 55",
            c3 and near(c3["shares"]["PSC"], .5 * .9) and near(c3["shares"]["F&F"], .5 * .1 + .5),
            c3 and c3["shares"])
        aos.remove(o3["id"], "acct", engine=eng)

        print("\n7. Only accounting writes; everyone reads")
        for n, role in (("acct1", "accountant"), ("ana", "analyst"), ("boss", "admin")):
            create_user(n, "pw", role=role)
        ids = {u["username"]: u["id"] for u in list_users()}

        def H(n, role):
            return {"Authorization": "Bearer " + jwt.encode(
                {"sub": str(ids[n]), "username": n, "role": role,
                 "exp": datetime.now(timezone.utc) + timedelta(hours=1)}, app.config["JWT_SECRET"], algorithm="HS256")}
        client = app.test_client()
        from flask_app.auth.sections import set_user_sections
        set_user_sections(ids["boss"], {"accounting": False}, "admin")   # Charlene's case
        body = {"entity_id": "PSC3", "investment_id": "DEALE", "effective_date": "2026-04-01",
                "lines": [{"investor_id": "OWPSC", "amount": 1}], "reason": "test"}
        for n, role in (("ana", "analyst"), ("boss", "admin")):
            r = client.post("/api/reports/pe-exposure/overrides", json=body, headers=H(n, role))
            chk("%s (%s%s) cannot add an override (403)" % (n, role, ", Accounting unticked" if n == "boss" else ""),
                r.status_code == 403, r.status_code)
            chk("...but reads them, told it cannot edit",
                client.get("/api/reports/pe-exposure/overrides", headers=H(n, role)).get_json()["can_edit"] is False)
        r = client.post("/api/reports/pe-exposure/overrides", json=body, headers=H("acct1", "accountant"))
        chk("an accountant with the Accounting section adds one (201)", r.status_code == 201, r.get_json())
        oid = (r.get_json() or {}).get("id")
        r = client.delete("/api/reports/pe-exposure/overrides/%s" % oid, headers=H("boss", "admin"))
        chk("the admin ROLE without Accounting cannot remove one (403)", r.status_code == 403, r.status_code)
        r = client.delete("/api/reports/pe-exposure/overrides/%s" % oid, headers=H("acct1", "accountant"))
        chk("the accountant removes it (200)", r.status_code == 200, r.status_code)
        r = client.get("/api/reports/pe-exposure/overrides/owners?entity=PSC3&as_of=2026-06-30",
                       headers=H("ana", "analyst"))
        chk("the investor list starts from MRI's commitments, with amounts",
            r.status_code == 200 and {o["investor_id"] for o in r.get_json()["owners"]} == {"OWPSC", "DCXVIA", "BRECO"}
            and all(o["committed"] for o in r.get_json()["owners"]), r.get_json())

        from database import PROTECTED_TABLES
        chk("both tables are protected from the CSV import (the app holds the only copy)",
            {"ownership_overrides", "ownership_override_lines"} <= PROTECTED_TABLES)
        raw = oc.group_shares("PPIC", Q2, pe.STOPS, context=pe._context, default_group=pe.DEFAULT_GROUP,
                              engine=eng, investment="DEALC", use_overrides=False)
        chk("use_overrides=False is the raw commitment walk (60/30/10)", near(raw["shares"]["PSC"], .6))

    print("\n%d passed, %d failed" % (len(_passed), len(_failed)))
    return 1 if _failed else 0


if __name__ == "__main__":
    sys.exit(main())
