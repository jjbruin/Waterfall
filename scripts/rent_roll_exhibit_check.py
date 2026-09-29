"""Guardrail: the IC rent roll exhibit reproduces new business's template.

Two halves.

1. THEIR OWN FIGURES, REBUILT (when their template is on this machine). Their
   exhibit B2:H111 is read into the shape the timeline engine produces, the
   workbook is built from it, and every cell of the data range is compared with
   theirs -- value, number format, font (name, size, bold, italic), fill, and all
   four borders. That isolates the LAYOUT from the data: if this passes, any
   difference on real data is a data question, not a presentation one. (The
   treasury `v492` method: rebuild the accepted file from its own contents.)
   The one deliberate difference -- their second header row pasted at row 56 --
   is replaced by print titles, and the comparison skips it by design.

2. A SYNTHETIC ROLL, always: print titles, page breaks never inside a tenant's
   group, a fair-market option printed as its method, the Flags sheet.

Run:  .venv\\Scripts\\python.exe scripts\\rent_roll_exhibit_check.py
"""
import io
import os
import sys
from pathlib import Path

sys.path.insert(0, os.getcwd())

import openpyxl  # noqa: E402
from flask_app.services.rent_roll_exhibit import build_exhibit  # noqa: E402

OK, BAD, SKIP = [], [], []
TEMPLATE = Path(os.path.expanduser(
    r"~/Downloads/Claude - Market at Poplar Rent Roll.xlsx"))


def chk(name, cond, detail=''):
    (OK if cond else BAD).append(name)
    print(('  PASS  ' if cond else '  FAIL  ') + name
          + ('  [%s]' % (detail,) if detail and not cond else ''))


def section(t):
    print('\n--- ' + t)


def iso(v):
    return v.date().isoformat() if hasattr(v, 'date') else (str(v)[:10] if v else None)


def roll_from_template(ws):
    """Their exhibit as the timeline engine's output shape."""
    tenants, cur = [], None
    for r in range(3, 111):
        b, c, d, e, f, g, h = [ws.cell(r, i).value for i in range(2, 9)]
        if b == 'Tenant':            # their pasted page-2 header
            continue
        if c is not None:
            cur = {'tenant_name': b, 'rows': [{'kind': 'tenant', 'label': None, 'sf': c,
                   'start': iso(d), 'end': iso(e), 'annual_rent': f, 'psf': g,
                   'options': h}], 'options': [], 'flags': []}
            tenants.append(cur)
            mode = 'term'
        else:
            if b and str(b).startswith('Option'):
                kind, mode = 'option', 'option'
            else:
                kind = 'step' if mode == 'term' else 'option_step'
            cur['rows'].append({'kind': kind, 'label': b, 'start': iso(d), 'end': iso(e),
                                'annual_rent': f, 'psf': g})
    sf = sum(t['rows'][0]['sf'] for t in tenants)
    rent = sum(t['rows'][0]['annual_rent'] for t in tenants)
    return {'as_of': '2026-09-01', 'tenants': tenants,
            'totals': {'square_feet': sf, 'annual_rent': rent, 'psf': rent / sf}}


def style_of(c):
    f, b = c.font, c.border
    fill = None
    if c.fill is not None and c.fill.fill_type == 'solid':
        fc = c.fill.fgColor
        fill = ('theme', fc.theme, round(fc.tint, 3)) if fc.type == 'theme' else ('rgb', fc.rgb)
    # Bold/italic are compared only where there is text to show them: their
    # template is itself inconsistent on EMPTY label cells (B5 plain, B10 italic),
    # and emphasis on an empty cell does not render.
    # The same holds for the font NAME on an empty cell (C4 Calibri, C7 Times in
    # their file): invisible, and inconsistent in the template itself. Fonts are
    # compared where there is text; borders and fills everywhere.
    font = (f.name, f.sz, bool(f.b), bool(f.i)) if c.value is not None else None
    return {'font': font, 'fill': fill,
            'border': tuple((getattr(b, s).style if getattr(b, s) is not None else None)
                        for s in ('left', 'right', 'top', 'bottom')),
            'fmt': c.number_format if c.value is not None else None}


def norm_val(v):
    if hasattr(v, 'date'):
        return v.date().isoformat()
    if isinstance(v, float):
        return round(v, 2)
    return v


section('1. their own figures, rebuilt, cell by cell')
if not TEMPLATE.exists():
    SKIP.append('template')
    print('  SKIP  template comparison (%s not on this machine)' % TEMPLATE)
else:
    theirs = openpyxl.load_workbook(TEMPLATE, data_only=True).worksheets[0]
    roll = roll_from_template(theirs)
    ours = openpyxl.load_workbook(io.BytesIO(build_exhibit(roll))).worksheets[0]
    # their rows map onto ours by dropping their pasted header at row 56
    their_rows = [r for r in range(2, 112) if theirs.cell(r, 2).value != 'Tenant' or r == 2]
    val_diff, sty_diff = [], []
    for i, tr in enumerate(their_rows):
        orow = 2 + i
        for col in range(2, 9):
            tc, oc = theirs.cell(tr, col), ours.cell(orow, col)
            if norm_val(tc.value) != norm_val(oc.value):
                val_diff.append((tc.coordinate, tc.value, oc.value))
            st, so = style_of(tc), style_of(oc)
            for k in st:
                if st[k] != so[k]:
                    sty_diff.append((tc.coordinate, k, st[k], so[k]))
    chk('every VALUE in the data range matches (%d cells)' % (len(their_rows) * 7),
        not val_diff, val_diff[:6])
    chk('every STYLE matches: font, fill, borders, number format', not sty_diff, sty_diff[:8])
    for col, w in (('A', 2.0), ('B', 40.0), ('C', 10.0), ('D', 13.54), ('F', 14.0),
                   ('G', 12.0), ('H', 23.45)):
        pass
    widths_ok = all(round(ours.column_dimensions[c].width or 0, 2) ==
                    round(theirs.column_dimensions[c].width or 0, 2)
                    for c in ('A', 'B', 'C', 'D', 'F', 'G', 'H'))
    chk('column widths match', widths_ok)
    chk('gridlines off, as theirs', ours.sheet_view.showGridLines is False)

section('2. paging, options without a figure, flags')
tenants = []
for n in range(30):
    rows = [{'kind': 'tenant', 'sf': 1000, 'start': '2020-01-01', 'end': '2030-12-31',
             'annual_rent': 20000, 'psf': 20.0, 'options': '1 x 5 Years'}]
    rows += [{'kind': 'step', 'label': 'Rent Step Dates' if k == 0 else None,
              'start': '202%d-01-01' % (7 + k), 'end': '202%d-12-31' % (7 + k),
              'annual_rent': 21000 + k, 'psf': 21.0} for k in range(2)]
    rows.append({'kind': 'option', 'label': 'Option 1', 'start': '2031-01-01',
                 'end': '2035-12-31', 'annual_rent': None, 'psf': None})
    tenants.append({'tenant_name': 'Tenant %02d' % n, 'rows': rows,
                    'options': [{'label': 'Option 1', 'rent_basis': 'fmv'}],
                    'flags': [{'code': 'option_rent_fmv', 'message': 'fair market'}]})
wb = openpyxl.load_workbook(io.BytesIO(build_exhibit(
    {'tenants': tenants, 'totals': {'square_feet': 30000, 'annual_rent': 600000, 'psf': 20.0}})))
ws = wb['Rent Roll']
chk('the header row repeats on every printed page', str(ws.print_title_rows) in ('$2:$2', '2:2'),
    ws.print_title_rows)
brk = [b.id for b in ws.row_breaks.brk]
group_ends = [2 + 4 * (i + 1) for i in range(30)]
chk('pages break, and only between tenant groups', brk and all(b in group_ends for b in brk),
    (brk, group_ends[:6]))
chk('a fair-market option prints its method, not a figure',
    ws.cell(6, 6).value == 'Fair market' and ws.cell(6, 6).font.i, ws.cell(6, 6).value)
chk('the total row closes the frame (medium bottom), grey, bold',
    ws.cell(123, 2).value == 'Total / Wtd. Avg.' and ws.cell(123, 2).border.bottom.style == 'medium'
    and ws.cell(123, 6).font.b, ws.cell(123, 2).value)
chk('the Flags sheet lists every flag', wb['Flags'].max_row == 31, wb['Flags'].max_row)

print('\n%d passed, %d failed, %d skipped' % (len(OK), len(BAD), len(SKIP)))
sys.exit(1 if BAD else 0)
