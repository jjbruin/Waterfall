"""Guardrail: asset management's Oct 6 2026 changes to the budget / Argus mapping.

Jack Day, Oct 6 2026:
  1. "Skip zero dollar lines. If a line has a $0 twelve month total, don't import it."
  2. "Export the mapping back to Excel ... drop it next to my budget file and pinpoint
     in a minute whether it's a sign flip or a line sitting in the wrong category."
  3. "Remove the checks/warnings section ... Keep the does-it-tie-to-the-spreadsheet
     check."
  4. Argus: "Add the flip sign checkbox. The budget upload has it and the Argus load
     doesn't."

What must stay true, asserted in BOTH directions where a rule could be satisfied by
doing too much:
  A. A line $0 in every month is set aside -- reported, not dropped -- on a fresh file
     AND a stored draft (mapping on it removed); a line that nets to $0 but moves months
     is KEPT; a line with no months read is KEPT (no evidence it is $0).
  B. ONE function for what a mapping writes: the budget commit, the Argus commit, the
     tie-out, the "as imported" figures and the export all go through
     `imported_amounts`, so the screen cannot show a figure the import does not write.
  C. Argus: the account sets the sign, the box (`reverse`) reverses it, and the
     budget's `flip` is IGNORED -- 14 production Argus lines carry an unseen `flip`.
  D. The export lists every line in sheet order with its sheet row, account, flip,
     sheet total and imported total; $0 lines appear as not imported; the by-account
     and tie-out sheets agree with the import's own figures.
  E. The screen shows no warnings, keeps the tie-out, offers the export, and offers
     the flip box for both sources.

Usage: python scripts/mapping_feedback_check.py [--inject=zero|netzero|argusflip|recon]
  zero      -- $0 lines not set aside
  netzero   -- the "months sum to $0" rule instead of "every month $0"
  argusflip -- Argus honours the budget's `flip`
  recon     -- the tie-out from the sheet total and `flip` (the old rule)
"""
import io
import os
import re
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

INJECT = next((a.split("=", 1)[1] for a in sys.argv[1:] if a.startswith("--inject=")), "")
_passed, _failed = [], []


def chk(label, cond, detail=""):
    (_passed if cond else _failed).append(label)
    print("   %s %s%s" % ("ok  " if cond else "FAIL", label,
                          ("   [%s]" % (detail,)) if detail and not cond else ""))


def main():
    os.environ.pop("DATABASE_URL", None)
    os.environ["DB_PATH"] = os.path.join(tempfile.mkdtemp(prefix="mapping_fb_"), "x.db")
    import openpyxl
    import sqlalchemy as sa
    from sqlalchemy import text
    from flask_app.services import budget_import_service as B
    from flask_app.services import budget_import_validate as V
    from flask_app.services import line_mapping_service as L
    from flask_app.services import valuation_service as VS

    if INJECT == "zero":
        B.without_zero_lines = lambda parsed, mapping=None: (
            parsed if mapping is None else (parsed, mapping))
        L.budget.without_zero_lines = B.without_zero_lines
    if INJECT == "netzero":
        _orig = B.without_zero_lines

        def _net(parsed, mapping=None):
            lines = parsed.get("lines") or []
            zero = [l for l in lines if round(sum((l.get("amounts") or {}).values()), 2) == 0
                    and l.get("amounts")]
            kept = [l for l in lines if l not in zero]
            p = {**parsed, "lines": kept,
                 "zero_lines": [{"row": l["row"], "label": l["label"]} for l in zero]}
            if mapping is None:
                return p
            gone = {str(l["row"]) for l in zero}
            return p, {k: v for k, v in mapping.items() if k not in gone}
        B.without_zero_lines = _net
        L.budget.without_zero_lines = _net
    if INJECT == "argusflip":
        V.flip_key = lambda source: "flip"
    if INJECT == "recon":
        _orig_recon = V.reconcile

        def _old_reconcile(parsed, mapping, source="budget"):
            # the pre-Oct-6 rule: sheet total x flip, whatever the source
            saved = V.imported_amounts
            V.imported_amounts = lambda line, m, src="budget": {
                "t": line["total"] * (-1 if m.get("flip") else 1)}
            try:
                return _orig_recon(parsed, mapping, source)
            finally:
                V.imported_amounts = saved
        V.reconcile = _old_reconcile

    # ── A. $0 lines ─────────────────────────────────────────────────────
    print("A. $0 lines are set aside, not dropped")
    wb = openpyxl.Workbook()
    ws = wb.active
    months = ["Jan 2027", "Feb 2027", "Mar 2027", "Apr 2027", "May 2027", "Jun 2027"]
    from datetime import date
    ws.append(["Account", *[date(2027, i + 1, 1) for i in range(6)]])
    ws.append(["4010 - Rental Income", 1000, 1000, 1000, 1000, 1000, 1000])
    ws.append(["5060 - Repair - Waterproofing", 0, 0, 0, 0, 0, 0])
    ws.append(["5060 - Repair - Roof", 500, 0, -500, 0, 0, 0])
    ws.append(["5050 - Utilities", 200, 200, 200, 200, 200, 200])
    buf = io.BytesIO()
    wb.save(buf)
    parsed = B.parse_budget_workbook(buf.getvalue(), "fixture.xlsx")
    labels = [l["label"] for l in parsed["lines"]]
    chk("the all-$0 line is not among the lines",
        not any("Waterproofing" in x for x in labels), labels)
    chk("...and is reported in zero_lines, with its sheet row",
        [z["label"] for z in parsed.get("zero_lines", [])] == ["5060 - Repair - Waterproofing"]
        and parsed["zero_lines"][0]["row"] == 2, parsed.get("zero_lines"))
    chk("a line that nets to $0 but moves months is KEPT",
        any("Roof" in x for x in labels), labels)
    chk("real lines are untouched", any("Rental" in x for x in labels)
        and any("Utilities" in x for x in labels), labels)

    # a stored draft: a $0 line and a mapping on it
    eng = sa.create_engine("sqlite:///:memory:")
    VS.ensure_valuation_tables(eng)
    with eng.begin() as c:
        cyc = c.execute(text("INSERT INTO valuation_cycles (year, as_of_date) "
                             "VALUES (2027, '2027-12-31') RETURNING id")).scalar()
        rid = c.execute(text("INSERT INTO valuation_records (cycle_id, vcode) "
                             "VALUES (:c, 'P0000018') RETURNING id"), {"c": cyc}).scalar()
    P = ["2027-01-31", "2027-02-28"]
    old = {"filename": "old.xlsx", "periods": P, "stated_totals": {}, "lines": [
        {"row": 5, "label": "Rent", "amounts": {P[0]: -1000.0, P[1]: -1000.0}, "total": -2000.0, "months": 2},
        {"row": 6, "label": "Waterproofing", "amounts": {P[0]: 0.0, P[1]: 0.0}, "total": 0.0, "months": 2},
        {"row": 7, "label": "No months read", "amounts": {}, "total": 0.0, "months": 0},
        {"row": 8, "label": "Repairs", "amounts": {P[0]: 300.0, P[1]: 200.0}, "total": 500.0, "months": 2},
    ]}
    omap = {"5": {"account": "4010", "flip": False}, "6": {"account": "5060", "flip": False},
            "8": {"account": "5060", "flip": True}}
    L.save_draft(eng, rid, "budget", "old.xlsx", old, omap, "jday")
    d = L.get_draft(eng, rid, "budget")
    chk("a stored draft loses its $0 line too",
        [l["label"] for l in d["parsed"]["lines"]] == ["Rent", "No months read", "Repairs"],
        [l["label"] for l in d["parsed"]["lines"]])
    chk("...and the mapping on it", "6" not in d["mapping"] and "5" in d["mapping"],
        sorted(d["mapping"]))
    chk("a line with NO months read is kept (no evidence it is $0)",
        any(l["label"] == "No months read" for l in d["parsed"]["lines"]))

    # ── B/C. one function for what is written ───────────────────────────
    print("B. One function for what a mapping writes")
    rent = {"row": 1, "label": "Rent", "amounts": {P[0]: 1000.0, P[1]: 1000.0}, "total": 2000.0}
    conc = {"row": 2, "label": "Free rent", "amounts": {P[0]: -300.0}, "total": -300.0}
    rep = {"row": 3, "label": "Repairs", "amounts": {P[0]: 250.0, P[1]: 0.0}, "total": 250.0}
    budget_w = V.imported_amounts(rent, {"account": "4010", "flip": True}, "budget")
    chk("budget: the sheet's figure, flipped when ticked",
        budget_w == {P[0]: -1000.0, P[1]: -1000.0}, budget_w)
    a = V.imported_amounts(rent, {"account": "4010"}, "argus")
    chk("Argus: revenue lands NEGATIVE in MRI's convention, from the account",
        a == {P[0]: -1000.0, P[1]: -1000.0}, a)
    a = V.imported_amounts(rent, {"account": "4010", "flip": True}, "argus")
    chk("Argus IGNORES the budget's flip (14 production lines carry one unseen)",
        a == {P[0]: -1000.0, P[1]: -1000.0}, a)
    a = V.imported_amounts(conc, {"account": "4010"}, "argus")
    chk("Argus: the account's rule alone makes a concession positive revenue",
        a == {P[0]: -300.0}, a)
    a = V.imported_amounts(conc, {"account": "4010", "reverse": True}, "argus")
    chk("...and the box reverses it", a == {P[0]: 300.0}, a)
    a = V.imported_amounts(rep, {"account": "5060"}, "argus")
    chk("Argus: $0 months are not written (as before)", a == {P[0]: 250.0}, a)

    src_v = (ROOT / "flask_app/services/budget_import_validate.py").read_text(encoding="utf-8")
    src_l = (ROOT / "flask_app/services/line_mapping_service.py").read_text(encoding="utf-8")
    commit_body = src_v.split("def commit(", 1)[1].split("\ndef ", 1)[0]
    argus_body = src_l.split("def _commit_argus(", 1)[1].split("\ndef ", 1)[0]
    export_body = src_l.split("def export_workbook(", 1)[1].split("\ndef ", 1)[0]
    chk("the budget commit writes imported_amounts",
        'imported_amounts(line, m, "budget")' in commit_body and "* flip" not in commit_body)
    chk("the Argus commit writes imported_amounts",
        'imported_amounts(line, m, "argus")' in argus_body
        and "_normalize_amount(coa, amount)" not in argus_body)
    chk("the export reads imported_amounts", "imported_amounts(line, m, source)" in export_body)

    tie_p = {"periods": P, "stated_totals": {}, "lines": [rent, conc, rep]}
    tie_m = {"1": {"account": "4010", "flip": True}, "2": {"account": "4010", "reverse": True},
             "3": {"account": "5060"}}
    rec = {r["line"]: r["computed"] for r in V.reconcile(tie_p, tie_m, "argus")["rows"]}
    chk("the Argus tie-out is what the Argus import writes (2,000 rent - 300 concession)",
        rec["revenue"] == 1700.0 and rec["expense"] == 250.0, rec)

    # ── D. the export ───────────────────────────────────────────────────
    print("D. The export")
    L.save_draft(eng, rid, "argus", "argus.xlsx", {**tie_p, "filename": "argus.xlsx",
                 "zero_lines": [{"row": 9, "label": "Zero line"}]}, tie_m, "jday")
    L.mark_draft_committed(eng, rid, "argus")
    content, fname = L.export_workbook(eng, rid, "argus", {"isbs_raw": None})
    xw = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    chk("three sheets: Lines, By account, Tie-out",
        xw.sheetnames == ["Lines", "By account", "Tie-out"], xw.sheetnames)
    ws = xw["Lines"]
    hdr = [c.value for c in ws[5]]
    rows = {r[1]: dict(zip(hdr, r)) for r in ws.iter_rows(min_row=6, values_only=True) if r[1]}
    chk("each line carries its SHEET row (parser row + 1)",
        rows["Rent"]["Sheet row"] == 2 and rows["Repairs"]["Sheet row"] == 4, rows.get("Rent"))
    chk("the flip column shows the Argus box, not the budget flip",
        rows["Rent"]["Flip sign"] == "No" and rows["Free rent"]["Flip sign"] == "Yes",
        (rows["Rent"]["Flip sign"], rows["Free rent"]["Flip sign"]))
    chk("total on the sheet beside total as imported",
        rows["Rent"]["Total on the spreadsheet"] == 2000.0
        and rows["Rent"]["Total as imported"] == -2000.0
        and rows["Free rent"]["Total as imported"] == 300.0, rows.get("Rent"))
    chk("the months are as imported", rows["Rent"][P[0]] == -1000.0, rows["Rent"].get(P[0]))
    chk("a $0 line is listed as not imported",
        "$0" in str(rows.get("Zero line", {}).get("Status")), rows.get("Zero line"))
    chk("the applied status is stated", "Applied" in str(ws["A2"].value), ws["A2"].value)
    ba = {r[0]: r for r in xw["By account"].iter_rows(min_row=2, values_only=True) if r[0]}
    chk("by account: 4010 is rent and concession combined, as imported",
        ba[4010][3] == 2 and ba[4010][4] == -1700.0, ba.get(4010))
    tie = {r[0]: r for r in xw["Tie-out"].iter_rows(min_row=2, max_row=4, values_only=True)}
    chk("the tie-out sheet is the tie-out", tie["Total revenue"][2] == 1700.0, tie)

    # ── E. the screen ───────────────────────────────────────────────────
    print("E. The screen")
    pn = (ROOT / "vue_app/src/components/common/LineMappingPanel.vue").read_text(encoding="utf-8")
    chk("no warnings are rendered", "criticalWarnings" not in pn and "infoWarnings" not in pn
        and "check.warnings" not in pn and "w.message" not in pn)
    chk("the tie-out is kept", "Does it tie to the spreadsheet?" in pn and "reconRows" in pn)
    chk("blocking problems still show (they stop the import)", "check.blocking" in pn)
    chk("the export is offered", "mapping/export" in pn and "Export mapping to Excel" in pn)
    chk("as imported is the server's figure, not re-derived",
        "check.value?.imported" in pn and "line.total || 0) * (m(line.row).flip" not in pn)
    chk("$0 lines are counted on screen", "parsed.zero_lines" in pn)
    api = (ROOT / "flask_app/api/valuations.py").read_text(encoding="utf-8")
    chk("the export route exists, signed-in only",
        re.search(r'mapping/export", methods=\["GET"\]\)\n@login_required', api) is not None)

    print("\n%d passed, %d failed" % (len(_passed), len(_failed)))
    return 1 if _failed else 0


if __name__ == "__main__":
    sys.exit(main())
