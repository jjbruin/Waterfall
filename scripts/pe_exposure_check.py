"""Guardrail: the PSC Preferred Equity exposure report (accounting's tracker).

Drives ``pe_exposure_service.build`` against a real SQLite database built here,
shaped like production. What it pins, each measured against accounting's 26Q2
tracker on Oct 5 2026:

  1. ONE ENGINE: each holding's capital balance IS the Pref Balance Detail
     engine's answer, and future funding IS the One Pager engine's. Asserted by
     calling both engines directly and comparing, not by re-deriving.
  2. Cost = balance + realized losses; a sale booked as a realized loss takes
     the holding to zero and off the report (Adirondack, City West).
  3. FMV = Cost + unrealized, INCLUDING a mark with no Effective Date (1,767
     non-cash rows carry none) and EXCLUDING one dated after the cut.
  4. The split walks commitments IN FORCE at the date, multiplied down: a JV at
     85/15, a TIAA JV at 90/10 whose 10% runs through an Ambassadors fund, PSC
     III to its outside holders, a holder that is itself a named group, a
     same-day tombstone ignored, and a revised commitment giving two quarters
     two different splits.
  5. Holders come from commitments, OP excluded; an InvestmentID shared by two
     vcodes is asked under the vcode the engines map it to.
  6. CAD converts at the stored Bank of Canada rate (USD = CAD / rate); with no
     rate the row is flagged and left out of the USD totals.
  7. Live (any date, not only a quarter end) and the Excel workbook.
"""
import io
import os
import sys
import tempfile
from datetime import date, datetime, timedelta, timezone

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

import pandas as pd  # noqa: E402
from sqlalchemy import text  # noqa: E402

_passed, _failed = [], []


def chk(label, cond, detail=None):
    (_passed if cond else _failed).append(label)
    print(("   ok   " if cond else "   FAIL ") + label + ("" if cond or detail is None else "   [%s]" % (detail,)))


def near(a, b, tol=0.01):
    return a is not None and b is not None and abs(float(a) - float(b)) <= tol


Q1, Q2 = date(2026, 3, 31), date(2026, 6, 30)

DEALS = [  # vcode, InvestmentID, name, currency
    ("P1", "DEALA", "Alpha Plaza", "USD"),
    ("P2", "DEALB", "Bravo Storage", "CAD"),
    ("P3", "DUP", "Dup Old", "USD"),
    ("P4", "DUP", "Dup Current", "USD"),
    ("P5", "SOLDX", "Sold Park", "USD"),
    ("P6", "DEALC", "Charlie Center", "USD"),
    ("P7", "DEALD", "Delta Direct", "USD"),
]

# EntityID, InvestorID, Amount, StartDate, EndDate
COMMITMENTS = [
    ("DEALA", "PPIA", 1000, "2020-01-01", None),
    ("DEALA", "OPALPHA", 500, "2020-01-01", None),          # an OP: never a holder
    ("PPIA", "PSC1", 300, "2020-01-01", "2026-03-31"),       # revised on 4/1 ...
    ("PPIA", "PSC1", 500, "2026-04-01", None),               # ... to 500
    ("PPIA", "PSCKOC", 700, "2020-01-01", None),
    ("PSCKOC", "KCREIT", 85, "2015-01-01", None),
    ("PSCKOC", "PSC1", 15, "2015-01-01", None),
    ("DEALB", "PPIB", 1421, "2021-01-01", None),
    ("PPIB", "TGAX", 100, "2021-01-01", None),
    ("TGAX", "TGAM", 90, "2021-01-01", None),
    ("TGAX", "INVX", 10, "2021-01-01", None),
    ("INVX", "AMBX", 100, "2021-01-01", None),
    ("AMBX", "PSC1", 74, "2021-01-01", None),
    ("AMBX", "OUTSIDER", 26, "2021-01-01", None),
    ("DUP", "PPIDUP", 500, "2021-01-01", None),
    ("SOLDX", "PPIS", 400, "2019-01-01", None),
    ("DEALC", "PPIC", 300, "2022-01-01", None),
    ("PPIC", "PSC3", 100, "2022-01-01", None),
    ("PPIC", "XTOMB", 0.01, "2026-06-30", "2026-06-30"),    # a same-day tombstone
    ("PSC3", "OWPSC", 60, "2016-01-01", None),
    ("PSC3", "DCXVIA", 30, "2016-01-01", None),
    ("PSC3", "BRECO", 10, "2016-01-01", None),
    ("DEALD", "PSS1", 100, "2022-01-01", None),              # a holder that IS a group
]

# The app's accounting feed: contributions NEGATIVE (the app's cashflow sign).
ACCT = [  # InvestmentID, InvestorID, EffectiveDate, MajorType, Typename, Amt, Capital
    ("DEALA", "PPIA", "2021-03-01", "Contribution", "Contribution: Investments", -1000, "Y"),
    ("DEALA", "PPIA", "2024-05-01", "Distribution", "Distribution: Return of Capital", 200, "Y"),
    ("DEALA", "PPIA", "2026-08-15", "Contribution", "Contribution: Investments", -50, "Y"),  # after Q2
    ("DEALA", "OPALPHA", "2021-03-01", "Contribution", "Contribution: Investments", -500, "Y"),
    ("DEALB", "PPIB", "2021-06-01", "Contribution", "Contribution: Investments", -1421, "Y"),
    ("DUP", "PPIDUP", "2021-06-01", "Contribution", "Contribution: Investments", -500, "Y"),
    ("SOLDX", "PPIS", "2019-06-01", "Contribution", "Contribution: Investments", -400, "Y"),
    ("DEALC", "PPIC", "2022-06-01", "Contribution", "Contribution: Investments", -300, "Y"),
    ("DEALD", "PSS1", "2022-06-01", "Contribution", "Contribution: Investments", -100, "Y"),
]

# accounting's IA query: accounting's signs. EffectiveDate None is real.
IA = [  # InvestmentID, InvestorID, TransactionDate, EffectiveDate, MajorType, Typename, Amount
    ("DEALA", "PPIA", "2025-12-31", None, "Other", "Unrealized Gain/Loss", 150.0),
    ("DEALA", "PPIA", "2026-09-30", "2026-09-30", "Other", "Unrealized Gain/Loss", 50.0),
    ("DEALA", "PPIA", "2025-06-30", "2025-06-30", "Other", "Realized Gain/Loss", 30.0),  # a GAIN: not counted
    ("SOLDX", "PPIS", "2025-12-31", "2025-12-31", "Other", "Realized Gain/Loss", -400.0),
    ("DEALB", "PPIB", "2025-12-31", "2025-12-31", "Other", "Unrealized Gain/Loss", 142.1),
]


def build_db(path):
    from sqlalchemy import create_engine
    eng = create_engine("sqlite:///" + path)
    pd.DataFrame(DEALS, columns=["vcode", "InvestmentID", "Investment_Name", "Currency"]).to_sql(
        "deals", eng, index=False, if_exists="replace")
    pd.DataFrame([dict(CommitmentUID=i, EntityID=e, InvestorID=v, Amount=a, CapitalPercent=0.0,
                       StartDate=s, EndDate=en) for i, (e, v, a, s, en) in enumerate(COMMITMENTS)]
                 ).to_sql("commitments", eng, index=False, if_exists="replace")
    pd.DataFrame(ACCT, columns=["InvestmentID", "InvestorID", "EffectiveDate", "MajorType", "Typename",
                                "Amt", "Capital"]).to_sql("accounting", eng, index=False, if_exists="replace")
    pd.DataFrame(IA, columns=["InvestmentID", "InvestorID", "TransactionDate", "EffectiveDate",
                              "MajorType", "Typename", "Amount"]).to_sql("ia_transactions", eng, index=False,
                                                                         if_exists="replace")
    pd.DataFrame(columns=["ENTITYID", "NAME"]).to_sql("entities", eng, index=False, if_exists="replace")
    pd.DataFrame(columns=["vcode", "vmisc", "iOrder", "PropCode", "dteffective", "mAmount", "nPercent",
                          "FXRate", "vState"]).to_sql("waterfalls", eng, index=False,
                                                                             if_exists="replace")
    from flask_app.services import market_rates_service as mr
    mr.ensure_tables(eng)
    with eng.begin() as c:
        c.execute(text("INSERT INTO market_rates VALUES ('USDCAD', '2026-06-30', 1.421, 'CAD per USD', "
                       "'Bank of Canada', '2026-07-01')"))
    return eng


def main() -> int:
    tmp = tempfile.mkdtemp(prefix="pe_exposure_check_")
    db = os.path.join(tmp, "check.db")
    os.environ.pop("DATABASE_URL", None)
    os.environ["DB_PATH"] = db
    eng = build_db(db)
    from flask_app import create_app
    app = create_app()
    app.config["DATABASE_URL"] = None
    from flask_app.services import pe_exposure_service as pe
    from flask_app.services.reports_service import build_pref_balance_detail
    data = {"acct": pd.read_sql("SELECT * FROM accounting", eng), "inv": pd.read_sql("SELECT * FROM deals", eng),
            "wf": pd.read_sql("SELECT * FROM waterfalls", eng),
            "commitments_raw": pd.read_sql("SELECT * FROM commitments", eng)}

    with app.app_context():
        r2 = pe.build(Q2, data=data, engine=eng)
        r1 = pe.build(Q1, data=data, engine=eng)
        live = pe.build(date(2026, 10, 5), data=data, engine=eng)
        r_norate = pe.build(date(2025, 12, 31), data=data, engine=eng)
        by = {(r["investment_id"], r["holder"]): r for r in r2["rows"]}

        print("\n1. One engine: the balance IS the Pref Balance Detail answer")
        a = by.get(("DEALA", "PPIA"))
        direct = build_pref_balance_detail("P1", "PPIA", Q2, data["acct"], data["inv"])["header"]["investment_balance"]
        chk("Alpha's balance equals build_pref_balance_detail's", a and near(a["balance"], direct), (a and a["balance"], direct))
        chk("...and is 800 (1,000 in, 200 back; the August call is after the cut)", a and near(a["balance"], 800))

        print("\n2. Cost = balance + realized losses; a realized gain is not counted")
        chk("Alpha's cost is 800 -- its +30 realized GAIN is ignored", a and near(a["cost"], 800) and near(a["realized_loss"], 0))
        chk("Sold Park (a sale booked as a -400 realized loss) is off the report",
            ("SOLDX", "PPIS") not in by, list(by))

        print("\n3. FMV = Cost + unrealized, through the cut")
        chk("Alpha FMV 950: the 150 mark with NO Effective Date is in", a and near(a["fmv"], 950), a and a["fmv"])
        chk("...the 50 mark dated 9/30 is not (it is in Live)",
            near({(r["investment_id"], r["holder"]): r for r in live["rows"]}[("DEALA", "PPIA")]["unrealized"], 200))

        print("\n4. The split: commitments in force on the date, multiplied down")
        s2, s1 = a["shares"], {(r["investment_id"], r["holder"]): r for r in r1["rows"]}[("DEALA", "PPIA")]["shares"]
        chk("Q1 (PSC1 300 / PSCKOC 700, KOC 85/15): PSC 40.5%, KOC 59.5%",
            near(s1["PSC"], 0.3 + 0.7 * 0.15, 1e-9) and near(s1["KOC"], 0.7 * 0.85, 1e-9), s1)
        chk("Q2 after the revision (500 / 700): a different split, PSC 50.4%",
            near(s2["PSC"], 500 / 1200 + 700 / 1200 * 0.15, 1e-9) and near(s2["KOC"], 700 / 1200 * 0.85, 1e-9), s2)
        b = by.get(("DEALB", "PPIB"))
        chk("TIAA JV 90/10 with its 10% through an Ambassadors fund: TIAA 90, PSC 7.4, Ambassadors 2.6",
            b and near(b["shares"]["TIAA"], 0.9, 1e-9) and near(b["shares"]["PSC"], 0.074, 1e-9)
            and near(b["shares"]["Ambassadors"], 0.026, 1e-9), b and b["shares"])
        c = by.get(("DEALC", "PPIC"))
        chk("PSC III to OWPSC 60 / Declaration 30 / Bruin (F&F) 10, the tombstone ignored",
            c and near(c["shares"]["PSC"], 0.6, 1e-9) and near(c["shares"]["Declaration"], 0.3, 1e-9)
            and near(c["shares"]["F&F"], 0.1, 1e-9) and not c["problems"], c and (c["shares"], c["problems"]))
        d = by.get(("DEALD", "PSS1"))
        chk("a holder that is itself PSC is PSC, whole", d and near(d["shares"]["PSC"], 1.0, 1e-9), d and d["shares"])
        chk("every resolvable row's shares total 100%",
            all(near(r["share_total"], 1.0, 1e-9) for r in r2["rows"] if not r["problems"]),
            [(r["holder"], r["share_total"]) for r in r2["rows"]])
        # PPIDUP has no owners on file. The split is NOT invented: the row says
        # so and the report's notes name it.
        dp = by.get(("DUP", "PPIDUP"))
        chk("a holder with no owners on file is reported, not given a made-up split",
            dp and dp["problems"] and dp["share_total"] == 0
            and any("PPIDUP" in n for n in r2["notes"]), dp and (dp["problems"], r2["notes"]))

        print("\n5. Holders from commitments; the engines' own vcode")
        chk("the OP is never a holder", not any(r["holder"] == "OPALPHA" for r in r2["rows"]))
        from loaders import build_investmentid_to_vcode
        want = str(build_investmentid_to_vcode(data["inv"])["DUP"]).strip()
        dup = [r for r in r2["rows"] if r["investment_id"] == "DUP"]
        chk("the shared InvestmentID appears once, under the engines' vcode with its 500",
            len(dup) == 1 and dup[0]["vcode"] == want and near(dup[0]["cost"], 500), [(x["vcode"], x["cost"]) for x in dup])

        print("\n6. CAD")
        chk("Bravo: 1,421 CAD / 1.421 = 1,000 USD; FMV 1,100", b and near(b["cost"], 1000) and near(b["fmv"], 1100)
            and b["fx"]["rate"] == 1.421, b and (b["cost"], b["fmv"], b["fx"]))
        nb = {(r["investment_id"], r["holder"]): r for r in r_norate["rows"]}.get(("DEALB", "PPIB"))
        chk("with no rate stored the CAD row is flagged, not converted", nb and nb["fx_missing"] and nb["cost"] is None,
            nb and (nb["fx_missing"], nb["cost"]))
        chk("...left out of the USD total, and the report says so",
            near(r_norate["totals"]["cost"], sum(x["cost"] or 0 for x in r_norate["rows"] if not x["fx_missing"]))
            and any("USD/CAD" in n for n in r_norate["notes"]), r_norate["notes"])
        chk("the Q2 total is the sum of its rows, CAD converted",
            near(r2["totals"]["cost"], sum(x["cost"] for x in r2["rows"])), r2["totals"]["cost"])

        print("\n7. Future funding is the One Pager engine's own figure")
        from one_pager import get_pe_performance
        pe_direct = get_pe_performance("P1", "2026-Q2", data["acct"][data["acct"].InvestmentID == "DEALA"],
                                       data["wf"], data["inv"], isbs_raw=None, deal_terms=None,
                                       commitments=data["commitments_raw"])
        ff = {f["vcode"]: f for f in r2["future_funding"]}.get("P1")
        rtf = pe_direct.get("remaining_to_fund")
        chk("Alpha's remaining to fund equals get_pe_performance's",
            (ff is None and (rtf is None or abs(rtf) < 1)) or (ff and near(ff["remaining_to_fund"], rtf)),
            (ff and ff["remaining_to_fund"], rtf))

        print("\n8. Live, and the workbook")
        chk("Live is dated today's date, not snapped to a quarter end",
            live["as_of"] == "2026-10-05" and live["is_quarter_end"] is False and live["future_funding_basis"])
        chk("...and Live's balance takes the August call (850)",
            near({(r["investment_id"], r["holder"]): r for r in live["rows"]}[("DEALA", "PPIA")]["balance"], 850))
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(pe.to_excel(r2)))
        chk("the workbook has the five sheets",
            wb.sheetnames == ["Cost", "FMV", "Future Funding Detail", "Ownership Routes", "Sources"], wb.sheetnames)
        ws = wb["Cost"]
        labels = [ws.cell(i, 1).value for i in range(1, ws.max_row + 1)]
        chk("future funding sits BELOW the current exposure, as accounting's tracker lays it out",
            labels.index("Net Invested Equity") < labels.index("Total Future Funding")
            < labels.index("Total Equity Invested / Committed") == ws.max_row - 1, labels)
        g = labels.index("Total Equity Invested / Committed") + 1
        n = labels.index("Net Invested Equity") + 1
        f = labels.index("Total Future Funding") + 1
        chk("the grand total adds the two subtotals", ws.cell(g, 7).value == f"=G{n}+G{f}", ws.cell(g, 7).value)
        chk("...and the report's own grand total is current exposure plus future funding",
            near(r2["totals"]["grand_cost"], r2["totals"]["cost"] + r2["totals"]["future_funding"])
            and near(r2["totals"]["grand_fmv"], r2["totals"]["fmv"] + r2["totals"]["future_funding"]))
        chk("the FX line names the rate and its source", "1.4210" in str(ws["A2"].value) and "Bank of Canada" in str(ws["A2"].value),
            ws["A2"].value)

    print("\n9. Open to the Reports section")
    from flask_app.auth import sections
    chk("/api/reports/pe-exposure is the Reports section's",
        sections.sections_for_api("/api/reports/pe-exposure") == ("reports",),
        sections.sections_for_api("/api/reports/pe-exposure"))

    print("\n%d passed, %d failed" % (len(_passed), len(_failed)))
    return 1 if _failed else 0


if __name__ == "__main__":
    sys.exit(main())
