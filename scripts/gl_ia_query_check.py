"""Guardrail: the CFO's GL / IA query tool.

His workbook `GL & IA Queries with Filters - 09182026.xlsx` lists what he needs to
vary. Every one of those is asserted here against a real database:

    GL   multiple entities / period / account(s) / export
    IA   multiple investment IDs / multiple investor IDs / date /
         MajorType(s) / SubType(s) / export

WHAT THIS IS NOT. It does not re-run his SQL. `queries/MRI_GL_Detail.sql` already IS
that query with the &SPARM smart parameters stripped out, and it imports into
`gl_detail`; the tool puts the parameters back against the copy we hold (Jim, Sep 19
2026: "since we are already pulling these tables into our database we can have the
query hit our tables"). A second copy of the SQL would be a second engine for the
same numbers.

THE CHECKS THAT MATTER MOST are the ones about what the tool says when it cannot
answer: an empty grid must be distinguishable from an unimported period, a truncated
grid must say it was truncated AND total the whole match rather than the visible
rows, and a filter value that is not on the allow-list must be refused rather than
reaching the SQL.

Run:  .venv\\Scripts\\python.exe scripts\\gl_ia_query_check.py
"""

import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from sqlalchemy import create_engine, text  # noqa: E402

from flask_app.services import gl_ia_query_service as Q  # noqa: E402

PASS, FAIL, SKIP = [], [], []


def check(name, cond, detail=''):
    (PASS if cond else FAIL).append(name)
    print(('  PASS  ' if cond else '  FAIL  ') + name + (f'  [{detail}]' if detail else ''))


def section(t):
    print(f'\n--- {t}')


# ---------------------------------------------------------------- fixture
DB = os.path.join(tempfile.gettempdir(), 'gl_ia_query_check.db')
if os.path.exists(DB):
    os.remove(DB)
eng = create_engine(f'sqlite:///{DB}')

with eng.begin() as c:
    c.execute(text("""CREATE TABLE gl_detail (
        "ENTITYID" TEXT, "PERIOD" TEXT, "ENTRDATE" TEXT, "ACCTNAME" TEXT,
        "ACCTNUM" TEXT, "BASIS" TEXT, "BALFOR" TEXT, "ITEM" INTEGER,
        "REF" TEXT, "DESCRPN" TEXT, "SEGMENTID" TEXT, "RLTDENTITY" TEXT,
        "RLTDENTITY_NAME" TEXT, "AMT" REAL)"""))
    c.execute(text("""CREATE TABLE ia_transactions (
        "InvestmentID" TEXT, "InvestmentName" TEXT, "InvestorID" TEXT,
        "InvestorName" TEXT, "TransactionDate" TEXT, "EffectiveDate" TEXT,
        "MajorType" TEXT, "Typename" TEXT, "SubtypeUID" TEXT, "Amount" REAL)"""))
    c.execute(text("""CREATE TABLE entities (
        "ENTITYID" TEXT, "NAME" TEXT, "ACTIVE" TEXT)"""))

    GL = [
        ('ENTA', '202401', '2024-01-15', 'Cash', '1000', 'A', 'B', 1, 'r', 'd', 's', '', '', 100.0),
        ('ENTA', '202402', '2024-02-15', 'Cash', '1000', 'B', 'N', 2, 'r', 'd', 's', '', '', 200.0),
        ('ENTA', '202402', '2024-02-20', 'Fees', '5000', 'B', 'N', 3, 'r', 'd', 's', '', '', -50.0),
        ('ENTB', '202401', '2024-01-31', 'Cash', '1000', 'A', 'B', 4, 'r', 'd', 's', '', '', 400.0),
        ('ENTB', '202403', '2024-03-31', 'Fees', '5000', 'B', 'N', 5, 'r', 'd', 's', '', '', -25.0),
    ]
    for row in GL:
        c.execute(text('INSERT INTO gl_detail VALUES (' +
                       ','.join(f':p{i}' for i in range(14)) + ')'),
                  {f'p{i}': v for i, v in enumerate(row)})
    for e, n in (('ENTA', 'Entity A LLC'), ('ENTB', 'Entity B LLC')):
        c.execute(text('INSERT INTO entities VALUES (:e, :n, :a)'),
                  {'e': e, 'n': n, 'a': 'Y'})

    IA = [
        ('INV1', 'Deal One', 'PPI01', 'PSC One', '2024-03-15', '2024-03-15',
         'Contribution', 'Capital Contribution', 'u1', 1000.0),
        ('INV1', 'Deal One', 'PPI02', 'PSC Two', '2024-06-30', '2024-06-30',
         'Distribution', 'Return of Capital', 'u2', -500.0),
        ('INV2', 'Deal Two', 'PPI01', 'PSC One', '2025-01-10', '2025-01-10',
         'Contribution', 'Capital Contribution', 'u1', 2000.0),
        ('INV2', 'Deal Two', 'PPI03', 'PSC Three', '2025-02-01', '2025-02-01',
         'Distribution', 'Preferred Return', 'u3', -75.0),
    ]
    for row in IA:
        c.execute(text('INSERT INTO ia_transactions VALUES (' +
                       ','.join(f':p{i}' for i in range(10)) + ')'),
                  {f'p{i}': v for i, v in enumerate(row)})


# ===========================================================================
section('The columns are the ones the CFO\'s queries return')

gl = Q.run_gl_query(eng)
ia = Q.run_ia_query(eng)
gl_labels = [c['label'] for c in gl['columns']]
ia_labels = [c['label'] for c in ia['columns']]

# His GL query selects ENTITYID, PERIOD, ENTRDATE, ACCTNAME, ACCTNUM, BASIS, BALFOR,
# ITEM, REF, DESCRPN, USER_DEFINED_VAL1, DEFINED_CODE, DESCRIPTION, AMT.
for lbl in ('Entity', 'Period', 'Entry Date', 'Account', 'Account Name', 'Basis',
            'Bal/Fwd', 'Item', 'Ref', 'Description', 'Segment', 'Related Entity',
            'Related Entity Name', 'Amount'):
    check(f'GL returns {lbl!r}', lbl in gl_labels)

# His IA query aliases them explicitly; these are his names.
for lbl in ('Investment ID', 'Investment Name', 'Investor ID', 'Investor Name',
            'Transaction Date', 'Effective Date', 'Major Type', 'Sub Type', 'Amount'):
    check(f'IA returns {lbl!r}', lbl in ia_labels)


# ===========================================================================
section('Every filter the CFO listed works, and narrows')

check('GL: unfiltered returns everything', gl['row_count'] == 5, str(gl['row_count']))

one = Q.run_gl_query(eng, entities=['ENTA'])
check('GL: one entity', one['row_count'] == 3, str(one['row_count']))
two = Q.run_gl_query(eng, entities=['ENTA', 'ENTB'])
check('GL: MULTIPLE entities, which is the actual ask',
      two['row_count'] == 5, str(two['row_count']))
check('...and two entities really is wider than one',
      two['row_count'] > one['row_count'])

acct = Q.run_gl_query(eng, accounts=['5000'])
check('GL: account(s)', acct['row_count'] == 2, str(acct['row_count']))
check('GL: multiple accounts',
      Q.run_gl_query(eng, accounts=['1000', '5000'])['row_count'] == 5)

per = Q.run_gl_query(eng, period_from='202402', period_to='202402')
check('GL: a single period', per['row_count'] == 2, str(per['row_count']))
check('GL: an open-ended period range',
      Q.run_gl_query(eng, period_from='202402')['row_count'] == 3)
check('GL: basis', Q.run_gl_query(eng, bases=['A'])['row_count'] == 2)

# Filters must AND, not OR -- ORing would quietly widen every query.
both = Q.run_gl_query(eng, entities=['ENTA'], accounts=['5000'])
check('GL: filters combine with AND', both['row_count'] == 1, str(both['row_count']))

# SEVERAL ENTITIES *AND* SEVERAL ACCOUNTS IN ONE QUERY (Jim, Sep 19 2026). One of
# each passing does not prove the pair: a bug that collapsed either list to its
# first element would still satisfy the single-value checks above.
multi = Q.run_gl_query(eng, entities=['ENTA', 'ENTB'], accounts=['1000', '5000'])
pairs = sorted({(r['ENTITYID'], r['ACCTNUM']) for r in multi['rows']})
check('GL: several entities AND several accounts together',
      pairs == [('ENTA', '1000'), ('ENTA', '5000'),
                ('ENTB', '1000'), ('ENTB', '5000')], str(pairs))
check('...and it is not silently using only the first entity',
      len({p[0] for p in pairs}) == 2)
check('...nor only the first account',
      len({p[1] for p in pairs}) == 2)
# Four of the five rows sit in 202401-202402; ENTB's 202403 row is the one the
# period drops, which is what makes this check about the period and not just the
# entity and account lists.
stacked = Q.run_gl_query(eng, entities=['ENTA', 'ENTB'], accounts=['1000', '5000'],
                         period_from='202401', period_to='202402',
                         bases=['A', 'B'])
check('...and every other filter still stacks on top of both',
      stacked['row_count'] == 4, str(stacked['row_count']))
check('...the period really dropped the out-of-range row',
      stacked['row_count'] < multi['row_count'],
      f"{stacked['row_count']} of {multi['row_count']}")

# The same on the IA side: several investments AND several investors at once.
iam = Q.run_ia_query(eng, investments=['INV1', 'INV2'],
                     investors=['PPI01', 'PPI02'])
ipairs = sorted({(r['InvestmentID'], r['InvestorID']) for r in iam['rows']})
check('IA: several investments AND several investors together',
      ipairs == [('INV1', 'PPI01'), ('INV1', 'PPI02'), ('INV2', 'PPI01')],
      str(ipairs))
check('...and PPI03 is excluded, so the investor list really applied',
      not any(r['InvestorID'] == 'PPI03' for r in iam['rows']))

check('IA: investment IDs', Q.run_ia_query(eng, investments=['INV1'])['row_count'] == 2)
check('IA: multiple investment IDs',
      Q.run_ia_query(eng, investments=['INV1', 'INV2'])['row_count'] == 4)
check('IA: investor IDs', Q.run_ia_query(eng, investors=['PPI01'])['row_count'] == 2)
check('IA: multiple investor IDs',
      Q.run_ia_query(eng, investors=['PPI01', 'PPI03'])['row_count'] == 3)
check('IA: major types',
      Q.run_ia_query(eng, major_types=['Contribution'])['row_count'] == 2)
check('IA: sub types',
      Q.run_ia_query(eng, sub_types=['Preferred Return'])['row_count'] == 1)
check('IA: a date range',
      Q.run_ia_query(eng, date_from='2025-01-01')['row_count'] == 2)
check('IA: the end date is INCLUSIVE, as the screen says it is',
      Q.run_ia_query(eng, date_to='2024-06-30')['row_count'] == 2,
      str(Q.run_ia_query(eng, date_to='2024-06-30')['row_count']))
check('IA: effective date can be used instead of transaction date',
      Q.run_ia_query(eng, date_field='EffectiveDate',
                     date_from='2025-01-01')['row_count'] == 2)


# ===========================================================================
section('Totals are over the whole match, not the visible page')

check('GL total is the sum of the matched rows',
      abs(gl['totals']['AMT'] - 625.0) < 1e-9, str(gl['totals']['AMT']))
check('IA total nets contributions against distributions',
      abs(ia['totals']['Amount'] - 2425.0) < 1e-9, str(ia['totals']['Amount']))

capped = Q.run_gl_query(eng, limit=2)
check('a capped result says how many rows there really are',
      capped['row_count'] == 5 and capped['shown'] == 2,
      f"{capped['shown']} of {capped['row_count']}")
check('...and flags itself as truncated', capped['truncated'] is True)
# The trap: totalling only the rows shown would make a truncated grid look complete
# AND wrong, and nothing on screen would say which.
check('...and STILL totals all five rows, not the two shown',
      abs(capped['totals']['AMT'] - 625.0) < 1e-9, str(capped['totals']['AMT']))
check('...and says so in a note the reader will see',
      any('not just the rows displayed' in n for n in capped['notes']),
      '; '.join(capped['notes']))
check('an untruncated result does not claim to be truncated',
      gl['truncated'] is False)


# ===========================================================================
section('What it cannot answer, it says')

early = Q.run_gl_query(eng, period_from='202001', period_to='202012')
check('a period before our import bound returns nothing', early['row_count'] == 0)
check('...and names the import bound as the reason, rather than looking empty',
      any(Q.GL_PERIOD_FLOOR in n for n in early['notes']), '; '.join(early['notes']))

inrange = Q.run_gl_query(eng, period_from='202402')
check('a period INSIDE the bound carries no such note',
      not any(Q.GL_PERIOD_FLOOR in n for n in inrange['notes']),
      '; '.join(inrange['notes']))

nomatch = Q.run_gl_query(eng, entities=['NOSUCH'])
check('a filter matching nothing returns an empty result, not an error',
      nomatch['available'] and nomatch['row_count'] == 0)

check('no completed MRI refresh reports freshness as unknown, not as now',
      gl['data_as_of'] is None)

# A missing table is a setup problem with a fix, not a crash.
empty_eng = create_engine('sqlite://')
miss = Q.run_gl_query(empty_eng)
check('a missing gl_detail is reported with what to do about it',
      miss['available'] is False and 'MRI_GL_Detail' in
      Q.gl_filter_options(empty_eng)['reason'])
check('...same for ia_transactions',
      Q.run_ia_query(empty_eng)['available'] is False)


# ===========================================================================
section('No filter value reaches the SQL as text')

# `date_field` is the one filter that names a COLUMN, so it is the one that could
# become an injection if it were interpolated. It is matched against a fixed set.
try:
    Q.run_ia_query(eng, date_field="TransactionDate; DROP TABLE ia_transactions--")
    check('an unknown date field is refused', False)
except ValueError as e:
    check('an unknown date field is refused', 'Unknown date field' in str(e))

with eng.connect() as c:
    still = c.execute(text('SELECT COUNT(*) FROM ia_transactions')).scalar()
check('...and the table is still there', still == 4, str(still))

# A value carrying SQL is data, not code: it simply matches nothing.
inj = Q.run_gl_query(eng, entities=["ENTA' OR '1'='1"])
check('a filter value carrying SQL matches nothing rather than everything',
      inj['row_count'] == 0, str(inj['row_count']))

# Blank and duplicate entries must not change the answer.
check('blank filter entries are ignored',
      Q.run_gl_query(eng, entities=['ENTA', '', '  '])['row_count'] == 3)
check('a repeated filter value does not duplicate rows',
      Q.run_gl_query(eng, entities=['ENTA', 'ENTA'])['row_count'] == 3)


# ===========================================================================
section('The pickers are built from the data, not from a hand-kept list')

opts = Q.gl_filter_options(eng)
check('GL options are available', opts['available'])
check('entities come with their names',
      {e['id']: e['name'] for e in opts['entities']}
      == {'ENTA': 'Entity A LLC', 'ENTB': 'Entity B LLC'})
check('periods are only those present', opts['periods'] == ['202401', '202402', '202403'])
check('accounts come with their names',
      [a['account'] for a in opts['accounts']] == ['1000', '5000'])
check('bases are only those present', sorted(opts['bases']) == ['A', 'B'])

iopts = Q.ia_filter_options(eng)
check('IA options are available', iopts['available'])
check('investments come with their names',
      [i['id'] for i in iopts['investments']] == ['INV1', 'INV2'])
check('investors come with their names',
      [i['id'] for i in iopts['investors']] == ['PPI01', 'PPI02', 'PPI03'])
check('major types are listed', sorted(iopts['major_types']) ==
      ['Contribution', 'Distribution'])
# A sub type belongs to a major type. A flat list would let "Return of Capital"
# under Distribution be picked as though it were the Contribution one.
check('sub types carry the major type they belong to',
      all('major_type' in s and 'sub_type' in s for s in iopts['sub_types']))
check('...and the pairing is right',
      {'major_type': 'Distribution', 'sub_type': 'Return of Capital'}
      in iopts['sub_types'])


# ===========================================================================
section('Export is a real workbook that records what produced it')

xl = Q.to_excel(Q.run_gl_query(eng, entities=['ENTA']), 'GL Detail',
                ['Entities: ENTA', 'Rows: 3'])
check('the export is a valid xlsx', xl[:2] == b'PK', str(xl[:2]))

import io as _io  # noqa: E402
import openpyxl  # noqa: E402
wb = openpyxl.load_workbook(_io.BytesIO(xl))
ws = wb.active
cells = [str(ws.cell(r, 1).value) for r in range(1, 12)]
check('...names the sheet', ws.title == 'GL Detail', ws.title)
# An exported grid with no record of the filters cannot be checked or repeated, and
# these get mailed around.
check('...records the filters that produced it',
      any('Entities: ENTA' in c for c in cells), str(cells[:5]))
check('...and carries the rows',
      ws.max_row >= 3 + 3, f'max_row={ws.max_row}')

vals = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
check('...only the filtered entity is in it',
      'ENTB' not in [str(v) for v in vals])

check('the export cap is higher than the screen cap, since exporting IS the way '
      'to get every row', Q.EXPORT_MAX_ROWS > Q.MAX_ROWS,
      f'{Q.EXPORT_MAX_ROWS} vs {Q.MAX_ROWS}')


# ===========================================================================
section('The endpoints exist and are reachable')

from flask_app import create_app  # noqa: E402

app = create_app()
rules = {str(r.rule) for r in app.url_map.iter_rules()}
for path in ('/api/gl-ia-query/gl', '/api/gl-ia-query/gl/options',
             '/api/gl-ia-query/gl/excel', '/api/gl-ia-query/ia',
             '/api/gl-ia-query/ia/options', '/api/gl-ia-query/ia/excel'):
    check(f'{path} is registered', path in rules)

_view = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                     'vue_app', 'src', 'components', 'layout', 'AppSidebar.vue')
if os.path.exists(_view):
    _src = open(_view, encoding='utf-8').read()
    check('the tool is in the sidebar', '/gl-ia-query' in _src)
    # Jim: "put the tool at the bottom of the list."
    check('...at the BOTTOM of Accounting, as asked',
          _src.index('to="/gl-ia-query"') > _src.index('to="/treasury"'))
    # Without this the section does not open when the route is reached directly.
    check('...and the section auto-expands on that route',
          "'/gl-ia-query'" in _src and 'acctRoutes' in _src)
else:
    SKIP.append('sidebar checks')
    print('  SKIP  sidebar checks  [Vue source not in the image]')


# ===========================================================================
section('The grid hides four columns; the EXPORT still carries them')
# Jim, Sep 19 2026: drop Bal/Fwd, Item, Related Entity and Related Entity Name
# from the screen so the key figures fit without scrolling, and narrow the
# description.
#
# ASSERTED IN BOTH DIRECTIONS ON PURPOSE. Hiding a column on screen is a
# display choice; dropping it from the workbook would lose the record somebody
# checks the screen against, and ITEM is how a line is found again in MRI's
# journal. A check written only for "these are hidden" is satisfied by deleting
# them outright, which is the one outcome that would do damage.
from flask_app.services import gl_ia_query_service as _gq  # noqa: E402

_ASKED = {'BALFOR', 'ITEM', 'RLTDENTITY', 'RLTDENTITY_NAME'}
check('exactly the four columns asked for are hidden',
      _gq.GL_SCREEN_HIDDEN == _ASKED, str(_gq.GL_SCREEN_HIDDEN))
_keys = [k for k, _ in _gq.GL_COLUMNS]
for _k in sorted(_ASKED):
    check(f'...{_k} is still SELECTED and still in the export',
          _k in _keys)
check('the money and the account are NOT hidden',
      not (_gq.GL_SCREEN_HIDDEN & {'AMT', 'ACCTNUM', 'ACCTNAME', 'ENTITYID',
                                   'PERIOD', 'ENTRDATE', 'DESCRPN'}))
check('the description is clipped, not hidden',
      'DESCRPN' in _gq.GL_CLIPPED and 'DESCRPN' not in _gq.GL_SCREEN_HIDDEN)

_res = _gq.run_gl_query(eng, limit=5)
if _res.get('available'):
    _cols = _res['columns']
    check('every column still comes back from the query',
          len(_cols) == len(_gq.GL_COLUMNS), f'{len(_cols)}')
    check('...each carrying whether the screen shows it',
          all('hidden' in c and 'clip' in c for c in _cols))
    _vis = [c['key'] for c in _cols if not c['hidden']]
    check('...and the four are flagged hidden',
          not (_ASKED & set(_vis)), str(sorted(set(_vis) & _ASKED)))
    check('...leaving the ten that matter', len(_vis) == len(_gq.GL_COLUMNS) - 4,
          str(len(_vis)))
    # The workbook is built from result['columns'], which is the FULL list.
    _xl = _gq.to_excel(_res, 'GL', ['check'])
    check('the workbook is still written', bool(_xl) and _xl[:2] == b'PK')
    import openpyxl as _op  # noqa: E402
    import io as _io  # noqa: E402
    _ws = _op.load_workbook(_io.BytesIO(_xl)).active
    _heads = {c.value for row in _ws.iter_rows() for c in row if c.value}
    for _lbl in ('Item', 'Bal/Fwd', 'Related Entity', 'Related Entity Name'):
        check(f'...with "{_lbl}" in it, hidden on screen or not', _lbl in _heads)
else:
    SKIP.append('grid column checks')
    print('  SKIP  grid column checks  [gl_detail not imported]')

_gview = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                      'vue_app', 'src', 'views', 'GlIaQueryView.vue')
if os.path.exists(_gview):
    _gsrc = open(_gview, encoding='utf-8').read()
    # A grid that still walks result.columns would render the hidden ones
    # anyway, and nothing on screen would say the flag was ignored.
    check('the grid walks the VISIBLE columns, not all of them',
          'v-for="c in shownCols"' in _gsrc
          and 'v-for="c in result.columns"' not in _gsrc)
    check('...filtering on the flag the server sends',
          '!c.hidden' in _gsrc)
    check('the clipped column carries its full value on hover',
          ':title="c.clip' in _gsrc)
else:
    SKIP.append('grid view checks')
    print('  SKIP  grid view checks  [Vue source not in the image]')




# ===========================================================================
section('The grid can be sorted and filtered by any column')
# Jim, Sep 20 2026: "take the query results that we are currently receiving and
# allow the user to filter or sort by any of the column headers of the query
# result." He asked for this after two candidate filters were measured and
# refused -- `ITEM = 1` (a line number, not a side) and the sign of `AMT` (which
# keeps the expense on an expense entry and the CASH on a revenue entry). Which
# line is the 'other side' is a property of the ACCOUNT, so the grid is made
# sliceable and the reader decides rather than a default deciding for them.
#
# Verified in the running app before these were written: sorting Amount gives
# -2,694,676.22 .. 2,614,646.68 ascending, reverses, and a third click restores
# the server's original order; filtering Account Name on "cash" takes 11 rows to
# 2 and the head reads "Total for all 11: (427.84)  Shown: 127,500.00".
_gview = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                      'vue_app', 'src', 'views', 'GlIaQueryView.vue')
if os.path.exists(_gview):
    _g = open(_gview, encoding='utf-8').read()

    # If the body still walks result.rows, every filter silently does nothing --
    # the boxes accept text, the count changes, the table does not.
    check('the table body walks the FILTERED rows',
          'v-for="(r, i) in viewRows"' in _g
          and 'v-for="(r, i) in result.rows"' not in _g)
    check('every visible column gets a filter box',
          'v-for="c in shownCols"' in _g and 'v-model="colFilters[c.key]"' in _g)
    check('every column header sorts', '@click="toggleSort(c.key)"' in _g)

    # Array.sort mutates. Sorting result.rows in place would destroy the
    # server's ordering permanently -- clearing the sort could not get it back.
    check('rows are COPIED before sorting, so clearing restores the original',
          'rows.slice().sort(' in _g)

    # The server totals the WHOLE match on purpose (a truncated grid that
    # totalled its own rows would look complete and be wrong). Once a filter is
    # on, that number no longer describes the screen, so both are shown.
    check('the whole-match total says how many rows it covers',
          'Total for all {{ result.row_count' in _g)
    check('...and a filtered subtotal is shown beside it, not instead of it',
          'viewTotal' in _g and 'Shown:' in _g)

    # Sorting 5,000 of 79,074 rows does not find the largest amount in the match.
    check('a truncated result says the slicing only covers what was loaded',
          'not to all' in _g and 'result.truncated' in _g)

    # A filter that matches nothing must say so rather than render an empty grid.
    check('filtering everything out is explained, not left blank',
          'No loaded row matches the column filters' in _g)

    # Numbers must sort as numbers: the cells render "(2,694,676.22)", so a
    # string sort would order by the bracket.
    check('numeric columns sort numerically, not as formatted text',
          "typeof x === 'number' && typeof y === 'number'" in _g)
    check('blanks sort to the end either way',
          'blanks last' in _g or 'return 1' in _g)
else:
    SKIP.append('grid slicing checks')
    print('  SKIP  grid slicing checks  [Vue source not in the image]')


print(f'\n{len(PASS)} passed, {len(FAIL)} failed, {len(SKIP)} skipped')
if FAIL:
    print('FAILED:')
    for f in FAIL:
        print('  - ' + f)
sys.exit(1 if FAIL else 0)
